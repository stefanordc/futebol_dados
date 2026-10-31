import argparse
import json
import random
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import pandas as pd
from bs4 import BeautifulSoup
from openpyxl.utils import get_column_letter
from playwright.sync_api import (
    Error as PWError,
    TimeoutError as PWTimeout,
    sync_playwright,
)


# ============================================================
# CONFIGURAÇÕES
# ============================================================

BASE_URL = "https://www.ogol.com.br"
COMPETICAO_URL = "https://www.ogol.com.br/competicao/liga-cazaque/256"

CAMPEONATO = "Cazaquistão"
PAIS = "Cazaque"


# ============================================================
# CAMINHOS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

CLUBES_PATH = BASE_DIR / "clubes.json"
CAMPEONATOS_PATH = BASE_DIR / "campeonatos.txt"
OUTPUT_PATH = Path(r"C:\Users\stefano.faria\Desktop\temporario.xlsx")


# ============================================================
# CONFIGURAÇÕES DO SCRAPER
# ============================================================

ROW_SELECTOR = "table.stats tbody tr[data-lj]"
EDICAO_SELECTOR = 'a[href^="/edicao/"]'

MAX_RETRIES = 3
MAX_PAGES = 100

PAGE_QUERY = "?fase_in=0&equipa=0&estado=&filtro=&op=calendario&page={pagina}"

# Imagens, fontes e mídia pesam na memória da aba e não são usadas no parsing
BLOQUEAR_RECURSOS = True
TIPOS_BLOQUEADOS = {"image", "font", "media"}


# ============================================================
# COLUNAS FINAIS
# ============================================================

COLUNAS_FINAIS = [
    "temporada",
    "pais",
    "campeonato",
    "jogo_id",
    "data",
    "rodada",
    "id_casa",
    "casa",
    "gols_casa",
    "gols_fora",
    "fora",
    "id_fora",
]


# ============================================================
# REGEX
# ============================================================

REGEX_DATA_ISO = re.compile(r"\b(\d{4}-\d{2}-\d{2})\b")
REGEX_DATA_BR = re.compile(r"\b(\d{2}/\d{2}/\d{4})\b")

# Aceita temporada de dois anos (2025/26, 2025/2026) e de ano único (2026)
REGEX_TEMPORADA = re.compile(r"\b(\d{4}(?:/\d{2,4})?)\b")

# Formatos de fase usados pelo site: R5, J5, 1/8, 1E, 2E, QF, SF, F, E...
REGEX_FASE = re.compile(
    r"(?:[RJ]\s*\d+|\d+\s*/\s*\d+|\d+\s*E|[A-Z]{1,3})",
)


# ============================================================
# NAVEGADOR
# ============================================================

USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
    "AppleWebKit/537.36 "
    "(KHTML, like Gecko) "
    "Chrome/154.0.0.0 Safari/537.36"
)

STEALTH_JS = """
Object.defineProperty(navigator, 'webdriver', {get: () => undefined});
Object.defineProperty(navigator, 'languages', {get: () => ['pt-BR', 'pt', 'en-US', 'en']});
Object.defineProperty(navigator, 'plugins', {get: () => [1, 2, 3, 4, 5]});
window.chrome = window.chrome || {runtime: {}};
"""


# ============================================================
# FUNÇÕES AUXILIARES
# ============================================================

def pausa(minimo=1.0, maximo=2.5):
    time.sleep(random.uniform(minimo, maximo))


def resumo_erro(erro):
    linhas = str(erro).strip().splitlines()

    return linhas[0] if linhas else type(erro).__name__


def filtrar_recurso(route):
    if route.request.resource_type in TIPOS_BLOQUEADOS:
        route.abort()
    else:
        route.continue_()


def criar_contexto(p, headless=True):
    browser = p.chromium.launch(
        headless=headless,
        args=[
            "--disable-blink-features=AutomationControlled",
            "--disable-dev-shm-usage",
        ],
    )

    context = browser.new_context(
        user_agent=USER_AGENT,
        locale="pt-BR",
        timezone_id="America/Sao_Paulo",
        viewport={
            "width": 1366,
            "height": 768,
        },
        extra_http_headers={
            "Accept-Language": "pt-BR,pt;q=0.9,en-US;q=0.8,en;q=0.7"
        },
    )

    context.add_init_script(STEALTH_JS)

    if BLOQUEAR_RECURSOS:
        context.route("**/*", filtrar_recurso)

    return browser, context


class Sessao:
    """
    Mantém o Playwright, o navegador e a aba. Se a aba travar ("Page crashed")
    ou o driver morrer, tudo é recriado do zero.
    """

    def __init__(self, headless=True):
        self.headless = headless
        self.p = None
        self.browser = None
        self.context = None
        self.page = None

        self.iniciar()

    def iniciar(self):
        self.p = sync_playwright().start()

        self.browser, self.context = criar_contexto(
            self.p,
            headless=self.headless,
        )

        self.page = self.context.new_page()

    def encerrar(self):
        for objeto in (self.context, self.browser):
            try:
                objeto.close()
            except Exception:
                pass

        try:
            self.p.stop()
        except Exception:
            pass

    def reiniciar(self):
        print(
            "   [!] Reiniciando o navegador...",
            file=sys.stderr,
        )

        self.encerrar()

        pausa(3, 6)

        self.iniciar()


def baixar_html(
    sessao,
    url,
    seletor,
    obrigatorio=True,
    timeout_seletor=30000,
):
    """
    Abre a URL e espera o seletor existir no DOM. Se a aba ou o driver
    caírem, reinicia o navegador e tenta de novo.
    """
    ultimo_erro = None

    for tentativa in range(1, MAX_RETRIES + 1):
        try:
            sessao.page.goto(
                url,
                wait_until="domcontentloaded",
                timeout=60000,
            )

            try:
                sessao.page.wait_for_selector(
                    seletor,
                    state="attached",
                    timeout=timeout_seletor,
                )
            except PWTimeout:
                if obrigatorio:
                    raise

            sessao.page.wait_for_timeout(800)

            return sessao.page.content()

        except PWTimeout as e:
            ultimo_erro = e

            print(
                f"[!] Tentativa {tentativa}/{MAX_RETRIES} falhou: {url}",
                file=sys.stderr,
            )

            pausa(3, 6)

        except PWError as e:
            ultimo_erro = e

            print(
                f"[!] Tentativa {tentativa}/{MAX_RETRIES} falhou "
                f"({resumo_erro(e)}): {url}",
                file=sys.stderr,
            )

            sessao.reiniciar()

    raise RuntimeError(
        f"Não foi possível carregar {url}. Erro: {resumo_erro(ultimo_erro)}"
    )


# ============================================================
# CLUBES.JSON
# ============================================================

def carregar_clubes():
    if not CLUBES_PATH.exists():
        sys.exit(
            f"[ERRO] clubes.json não encontrado: {CLUBES_PATH}"
        )

    try:
        with open(
            CLUBES_PATH,
            "r",
            encoding="utf-8-sig",
        ) as arquivo:
            dados = json.load(arquivo)

    except json.JSONDecodeError as e:
        sys.exit(
            f"[ERRO] clubes.json inválido: {e}"
        )

    clubes = {}

    for clube_id, nome in dados.items():
        try:
            chave = str(int(clube_id))
        except (ValueError, TypeError):
            chave = str(clube_id).strip()

        clubes[chave] = str(nome)

    print(
        f"clubes.json carregado: {len(clubes)} clubes"
    )

    return clubes


def nome_padronizado(clube_id, nome_site, clubes_json):
    if clube_id is None:
        return nome_site

    try:
        chave = str(int(clube_id))
    except (ValueError, TypeError):
        chave = str(clube_id).strip()

    if chave in clubes_json:
        return clubes_json[chave]

    return nome_site


# ============================================================
# TEMPORADAS
# ============================================================

def normalizar_temporada_texto(texto):
    if not texto:
        return None

    match = REGEX_TEMPORADA.search(texto)

    if not match:
        return None

    temporada = match.group(1)

    # Temporada de ano único (ex: MLS 2026)
    if "/" not in temporada:
        return temporada

    ano1, ano2 = temporada.split("/")

    if len(ano2) == 4:
        return f"{ano1}/{ano2}"

    ano1_int = int(ano1)
    ano2_int = int(ano2)

    seculo = (ano1_int // 100) * 100

    ano2_completo = seculo + ano2_int

    if ano2_completo < ano1_int:
        ano2_completo += 100

    return f"{ano1}/{ano2_completo}"


def listar_edicoes(html):
    soup = BeautifulSoup(
        html,
        "html.parser",
    )

    container = None

    for card in soup.select("div.card-data"):
        titulo = card.select_one("h2.header")

        if titulo is None:
            continue

        texto_titulo = titulo.get_text(
            " ",
            strip=True,
        ).lower()

        if "edições e vencedores" in texto_titulo:
            container = card
            break

        if "edicoes e vencedores" in texto_titulo:
            container = card
            break

    origem = container if container is not None else soup

    edicoes = []
    ids_encontrados = set()

    for a in origem.select(EDICAO_SELECTOR):
        href = a.get("href")

        if not href:
            continue

        href = href.strip()

        texto_card = a.get_text(
            " ",
            strip=True,
        )

        temporada = normalizar_temporada_texto(
            texto_card
        )

        if not temporada:
            continue

        match_id = re.search(
            r"/(\d+)/?$",
            href,
        )

        if not match_id:
            continue

        id_edicao = match_id.group(1)

        if id_edicao in ids_encontrados:
            continue

        ids_encontrados.add(id_edicao)

        url_edicao = urljoin(
            BASE_URL,
            href,
        ).rstrip("/")

        url_calendario = f"{url_edicao}/calendario"

        atual = "Atual" in texto_card

        edicoes.append(
            {
                "id_edicao": id_edicao,
                "temporada": temporada,
                "href": href,
                "url_edicao": url_edicao,
                "url_calendario": url_calendario,
                "atual": atual,
            }
        )

    edicoes.sort(
        key=lambda x: int(
            x["temporada"].split("/")[0]
        ),
        reverse=True,
    )

    return edicoes


# ============================================================
# CAMPEONATOS.TXT
# ============================================================

def salvar_link_temporada_atual(
    temporada,
    url_calendario,
):
    nova_linha = (
        f"{CAMPEONATO}|{temporada}|{url_calendario}"
    )

    linhas = []

    if CAMPEONATOS_PATH.exists():
        with open(
            CAMPEONATOS_PATH,
            "r",
            encoding="utf-8",
        ) as arquivo:
            linhas = [
                linha.strip()
                for linha in arquivo
                if linha.strip()
            ]

    prefixo = f"{CAMPEONATO}|"

    linhas = [
        linha
        for linha in linhas
        if not linha.startswith(prefixo)
    ]

    linhas.append(nova_linha)

    CAMPEONATOS_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with open(
        CAMPEONATOS_PATH,
        "w",
        encoding="utf-8",
    ) as arquivo:
        arquivo.write(
            "\n".join(linhas) + "\n"
        )

    print(
        f"Temporada atual: {temporada}"
    )

    print(
        f"Calendário atual: {url_calendario}"
    )


# ============================================================
# PARSING
# ============================================================

def parse_placar(texto):
    if not texto:
        return None, None

    match = re.search(
        r"(?<!\d)(\d+)\s*-\s*(\d+)(?!\d)",
        texto,
    )

    if not match:
        return None, None

    return (
        int(match.group(1)),
        int(match.group(2)),
    )


def limpar_fase(texto):
    """
    Devolve a fase exatamente como aparece no site (1/8, 1E, 2E, QF, E, R5...),
    apenas normalizando espaços.
    """
    if not texto:
        return None

    texto = re.sub(
        r"\s+",
        " ",
        str(texto),
    ).strip()

    return texto or None


def titulo_pagina(html):
    """Título (h1) da página, para conferir a qual competição a edição pertence."""
    soup = BeautifulSoup(
        html,
        "html.parser",
    )

    h1 = soup.select_one("h1")

    if h1 is None:
        return None

    return h1.get_text(
        " ",
        strip=True,
    ) or None


def extrair_data_linha(tr):
    candidatos = []

    td_data = tr.select_one("td.date")

    if td_data:
        candidatos.append(
            td_data.get_text(
                " ",
                strip=True,
            )
        )

    for td in tr.find_all("td"):
        candidatos.append(
            td.get_text(
                " ",
                strip=True,
            )
        )

    candidatos.append(
        tr.get_text(
            " ",
            strip=True,
        )
    )

    for texto in candidatos:
        match = REGEX_DATA_ISO.search(
            texto
        )

        if match:
            return match.group(1)

        match = REGEX_DATA_BR.search(
            texto
        )

        if match:
            return match.group(1)

    for elemento in tr.find_all(True):
        for atributo in (
            "data-date",
            "datetime",
            "title",
        ):
            valor = elemento.get(
                atributo
            )

            if not valor:
                continue

            match = REGEX_DATA_ISO.search(
                str(valor)
            )

            if match:
                return match.group(1)

            match = REGEX_DATA_BR.search(
                str(valor)
            )

            if match:
                return match.group(1)

    return None


def extrair_rodada_linha(tr):
    seletores = [
        "td.phase",
        "td.round",
        "td[class*='phase']",
        "td[class*='round']",
    ]

    for seletor in seletores:
        elemento = tr.select_one(
            seletor
        )

        if elemento:
            fase = limpar_fase(
                elemento.get_text(
                    " ",
                    strip=True,
                )
            )

            if fase is not None:
                return fase

    for td in tr.find_all("td"):
        texto = td.get_text(
            " ",
            strip=True,
        )

        if REGEX_FASE.fullmatch(texto):
            return limpar_fase(
                texto
            )

        if re.fullmatch(
            r"\d+\s*(?:ª|º)?\s*rodada",
            texto,
            flags=re.IGNORECASE,
        ):
            return limpar_fase(
                texto
            )

    return None


def extrair_jogo_id(tr):
    candidatos = [
        tr.get("id"),
        tr.get("data-lj"),
    ]

    for candidato in candidatos:
        if candidato is None:
            continue

        texto = str(
            candidato
        ).strip()

        if texto.isdigit():
            return int(texto)

        match = re.search(
            r"(\d+)",
            texto,
        )

        if match:
            return int(match.group(1))

    return None


def id_da_logo(td):
    if td is None:
        return None

    img = td.select_one("img")

    if img is None:
        return None

    src = img.get(
        "src",
        "",
    )

    match = re.search(
        r"/equipas/(\d+)_imgbank",
        src,
    )

    if not match:
        return None

    return int(match.group(1))


def id_do_link(a):
    if a is None:
        return None

    href = a.get(
        "href",
        "",
    )

    padroes = [
        r"/equipe/[^/?]+/(\d+)",
        r"/equipa/[^/?]+/(\d+)",
    ]

    for padrao in padroes:
        match = re.search(
            padrao,
            href,
        )

        if match:
            return int(match.group(1))

    return None


def ids_do_h2h(a):
    if a is None:
        return set()

    href = a.get(
        "href",
        "",
    )

    match = re.search(
        r"/t(\d+)-t(\d+)",
        href,
    )

    if not match:
        return set()

    return {
        int(match.group(1)),
        int(match.group(2)),
    }


# ============================================================
# EXTRAÇÃO DE UMA PÁGINA
# ============================================================

def parse_pagina(
    html,
    clubes_json,
):
    """
    Devolve (jogos finalizados, maior página citada, ids de TODAS as linhas).
    Os ids de todas as linhas (com ou sem placar) servem para detectar
    páginas repetidas e encerrar a paginação.
    """
    soup = BeautifulSoup(
        html,
        "html.parser",
    )

    jogos = []
    ids_linhas = set()

    for tr in soup.select(
        ROW_SELECTOR
    ):
        jogo_id = extrair_jogo_id(
            tr
        )

        if jogo_id is not None:
            ids_linhas.add(
                jogo_id
            )

        data = extrair_data_linha(
            tr
        )

        rodada = extrair_rodada_linha(
            tr
        )

        td_casa = tr.select_one(
            "td.text.home"
        )

        td_fora = tr.select_one(
            "td.text.away"
        )

        casa_link = (
            td_casa.select_one("a")
            if td_casa
            else None
        )

        fora_link = (
            td_fora.select_one("a")
            if td_fora
            else None
        )

        td_resultado = tr.select_one(
            "td.result"
        )

        texto_resultado = (
            td_resultado.get_text(
                " ",
                strip=True,
            )
            if td_resultado
            else None
        )

        gols_casa, gols_fora = parse_placar(
            texto_resultado
        )

        # JOGO SEM PLACAR NÃO É SALVO
        if (
            gols_casa is None
            or gols_fora is None
        ):
            continue

        id_casa = None
        id_fora = None

        if td_casa:
            td_logo_casa = (
                td_casa.find_next_sibling(
                    "td"
                )
            )

            id_casa = id_da_logo(
                td_logo_casa
            )

        if td_fora:
            td_logo_fora = (
                td_fora.find_previous_sibling(
                    "td"
                )
            )

            id_fora = id_da_logo(
                td_logo_fora
            )

        if id_casa is None:
            id_casa = id_do_link(
                casa_link
            )

        if id_fora is None:
            id_fora = id_do_link(
                fora_link
            )

        h2h = tr.select_one(
            "td.h2h a"
        )

        ids_h2h = ids_do_h2h(
            h2h
        )

        if len(ids_h2h) == 2:
            if (
                id_casa is None
                and id_fora is not None
            ):
                restantes = (
                    ids_h2h
                    - {id_fora}
                )

                if len(restantes) == 1:
                    id_casa = next(
                        iter(restantes)
                    )

            elif (
                id_fora is None
                and id_casa is not None
            ):
                restantes = (
                    ids_h2h
                    - {id_casa}
                )

                if len(restantes) == 1:
                    id_fora = next(
                        iter(restantes)
                    )

        ids_conferem = None

        if (
            len(ids_h2h) == 2
            and id_casa is not None
            and id_fora is not None
        ):
            ids_conferem = (
                {id_casa, id_fora}
                == ids_h2h
            )

        nome_casa_site = (
            casa_link.get_text(
                " ",
                strip=True,
            )
            if casa_link
            else None
        )

        nome_fora_site = (
            fora_link.get_text(
                " ",
                strip=True,
            )
            if fora_link
            else None
        )

        nome_casa = nome_padronizado(
            id_casa,
            nome_casa_site,
            clubes_json,
        )

        nome_fora = nome_padronizado(
            id_fora,
            nome_fora_site,
            clubes_json,
        )

        jogos.append(
            {
                "jogo_id": jogo_id,
                "data": data,
                "rodada": rodada,
                "id_casa": id_casa,
                "casa": nome_casa,
                "gols_casa": gols_casa,
                "gols_fora": gols_fora,
                "fora": nome_fora,
                "id_fora": id_fora,
                "_ids_conferem": ids_conferem,
            }
        )

    max_pagina = 1

    for a in soup.select(
        'a[href*="page="]'
    ):
        href = a.get(
            "href",
            "",
        )

        match = re.search(
            r"[?&]page=(\d+)",
            href,
        )

        if match:
            pagina = int(
                match.group(1)
            )

            max_pagina = max(
                max_pagina,
                pagina,
            )

    return jogos, max_pagina, ids_linhas


# ============================================================
# COLETA DE UMA TEMPORADA
# ============================================================

def raspar_temporada(
    sessao,
    url_calendario,
    temporada,
    clubes_json,
):
    todos = []
    vistos = set()

    ultima_pagina = 1
    pagina = 1
    paginas_lidas = 0

    while (
        pagina <= ultima_pagina
        and pagina <= MAX_PAGES
    ):
        if pagina == 1:
            url = url_calendario
        else:
            url = (
                url_calendario
                + PAGE_QUERY.format(
                    pagina=pagina
                )
            )

        html = baixar_html(
            sessao,
            url,
            "table.stats",
            obrigatorio=False,
            timeout_seletor=20000,
        )

        if pagina == 1:
            titulo = titulo_pagina(
                html
            )

            if titulo:
                print(
                    f"   Título: {titulo}"
                )

        jogos, max_pagina, ids_linhas = parse_pagina(
            html,
            clubes_json,
        )

        # Página sem nenhuma linha de jogo: acabou
        if not ids_linhas:
            break

        # Página que não traz nenhum jogo novo: o site está repetindo
        # conteúdo além da última página real
        if (
            pagina > 1
            and not (ids_linhas - vistos)
        ):
            print(
                f"   Página {pagina} repetida; encerrando paginação."
            )
            break

        vistos |= ids_linhas
        paginas_lidas += 1

        ultima_pagina = max(
            ultima_pagina,
            max_pagina,
        )

        todos.extend(
            jogos
        )

        pagina += 1

        if pagina <= ultima_pagina:
            pausa()

    if pagina > MAX_PAGES:
        print(
            f"   [!] Limite de {MAX_PAGES} páginas atingido.",
            file=sys.stderr,
        )

    print(
        f"   Páginas lidas: {paginas_lidas} | "
        f"linhas de jogo: {len(vistos)}"
    )

    if not todos:
        return pd.DataFrame(
            columns=[
                "jogo_id",
                "data",
                "temporada",
                "campeonato",
                "pais",
                "rodada",
                "id_casa",
                "casa",
                "gols_casa",
                "gols_fora",
                "fora",
                "id_fora",
                "_ids_conferem",
            ]
        )

    df = pd.DataFrame(
        todos
    )

    df["temporada"] = temporada
    df["campeonato"] = CAMPEONATO
    df["pais"] = PAIS

    com_id = df[
        df["jogo_id"].notna()
    ].copy()

    sem_id = df[
        df["jogo_id"].isna()
    ].copy()

    com_id = (
        com_id
        .drop_duplicates(
            subset=[
                "jogo_id",
                "temporada",
                "campeonato",
            ],
            keep="last",
        )
    )

    df = pd.concat(
        [
            com_id,
            sem_id,
        ],
        ignore_index=True,
    )

    return df


# ============================================================
# VALIDAÇÃO
# ============================================================

def validar_temporada(
    df,
    temporada,
):
    if df.empty:
        print(
            f"   -> {temporada}: nenhum jogo finalizado"
        )
        return

    problemas = []

    sem_jogo_id = int(
        df["jogo_id"].isna().sum()
    )

    sem_data = int(
        df["data"].isna().sum()
    )

    sem_rodada = int(
        df["rodada"].isna().sum()
    )

    sem_id_clube = int(
        (
            df["id_casa"].isna()
            | df["id_fora"].isna()
        ).sum()
    )

    divergentes = int(
        (
            df["_ids_conferem"]
            == False
        ).sum()
    )

    if sem_jogo_id:
        problemas.append(
            f"{sem_jogo_id} sem jogo_id"
        )

    if sem_data:
        problemas.append(
            f"{sem_data} sem data"
        )

    if sem_rodada:
        problemas.append(
            f"{sem_rodada} sem rodada"
        )

    if sem_id_clube:
        problemas.append(
            f"{sem_id_clube} sem ID de clube"
        )

    if divergentes:
        problemas.append(
            f"{divergentes} com divergência de IDs"
        )

    status = (
        "OK"
        if not problemas
        else "ATENÇÃO"
    )

    print(
        f"   -> {temporada}: "
        f"{len(df)} jogos finalizados [{status}]"
    )

    for problema in problemas:
        print(
            f"      {problema}"
        )

    if sem_data:
        print(
            "      Exemplos sem data:"
        )

        for _, linha in (
            df[
                df["data"].isna()
            ]
            .head(10)
            .iterrows()
        ):
            print(
                f"      jogo_id={linha['jogo_id']} | "
                f"{linha['casa']} x {linha['fora']}"
            )

    if sem_rodada:
        print(
            "      Exemplos sem rodada:"
        )

        for _, linha in (
            df[
                df["rodada"].isna()
            ]
            .head(10)
            .iterrows()
        ):
            print(
                f"      jogo_id={linha['jogo_id']} | "
                f"{linha['casa']} x {linha['fora']}"
            )


# ============================================================
# DATA
# ============================================================

def converter_data(valor):
    if valor is None:
        return None

    if pd.isna(valor):
        return None

    texto = str(
        valor
    ).strip()

    formatos = [
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
    ]

    for formato in formatos:
        try:
            return datetime.strptime(
                texto,
                formato,
            ).date()

        except ValueError:
            continue

    return None


# ============================================================
# FINAL x GRUPO F
# ============================================================

def marcar_final(df):
    """
    O site usa "F" tanto para o Grupo F quanto para a Final. Em cada temporada,
    a Final é o jogo "F" disputado depois de todos os jogos das demais fases;
    os jogos "F" da fase de grupos acontecem antes e continuam como "F".
    """
    for _, grupo in df.groupby("temporada"):
        outras_fases = grupo[
            grupo["rodada"].notna()
            & (grupo["rodada"] != "F")
            & grupo["data"].notna()
        ]

        # Sem outras fases para comparar, não dá para afirmar que é final
        if outras_fases.empty:
            continue

        ultima_data = outras_fases["data"].max()

        mascara = (
            (grupo["rodada"] == "F")
            & grupo["data"].notna()
            & (grupo["data"] >= ultima_data)
        )

        df.loc[grupo.index[mascara], "rodada"] = "Final"

    return df


# ============================================================
# MONTAR SAÍDA
# ============================================================

def montar_saida(partes):
    partes_validas = [
        df
        for df in partes
        if not df.empty
    ]

    if not partes_validas:
        return pd.DataFrame(
            columns=COLUNAS_FINAIS
        )

    df = pd.concat(
        partes_validas,
        ignore_index=True,
    )

    df = df[
        df["gols_casa"].notna()
        & df["gols_fora"].notna()
    ].copy()

    df["data"] = df[
        "data"
    ].apply(
        converter_data
    )

    for coluna in (
        "jogo_id",
        "id_casa",
        "id_fora",
        "gols_casa",
        "gols_fora",
    ):
        df[coluna] = pd.to_numeric(
            df[coluna],
            errors="coerce",
        ).astype("Int64")

    df = df[
        df["jogo_id"].notna()
    ].copy()

    df = (
        df
        .drop_duplicates(
            subset=[
                "jogo_id",
                "campeonato",
                "temporada",
            ],
            keep="last",
        )
        .reset_index(
            drop=True
        )
    )

    df = marcar_final(
        df
    )

    df["_ano"] = (
        df["temporada"]
        .astype(str)
        .str[:4]
    )

    df["_ano"] = pd.to_numeric(
        df["_ano"],
        errors="coerce",
    )

    # A fase é texto livre (1/8, 1E, QF...), então não entra na ordenação
    df = df.sort_values(
        by=[
            "_ano",
            "data",
            "jogo_id",
        ],
        ascending=[
            True,
            True,
            True,
        ],
        na_position="last",
    )

    df = df.drop(
        columns=[
            "_ano",
        ]
    )

    return df[
        COLUNAS_FINAIS
    ].reset_index(
        drop=True
    )


# ============================================================
# EXCEL
# ============================================================

def ajustar_excel(writer):
    ws = writer.sheets[
        "calendario"
    ]

    cabecalhos = [
        celula.value
        for celula in ws[1]
    ]

    coluna_data = get_column_letter(
        cabecalhos.index("data") + 1
    )

    for celula in ws[coluna_data][1:]:
        if celula.value:
            celula.number_format = (
                "dd/mm/yyyy"
            )

    # Força a fase como texto para o Excel não converter "1/8" em data
    coluna_rodada = get_column_letter(
        cabecalhos.index("rodada") + 1
    )

    for celula in ws[coluna_rodada][1:]:
        if celula.value is not None:
            celula.number_format = "@"

    for coluna in ws.columns:
        indice = coluna[0].column

        letra = get_column_letter(
            indice
        )

        maior = 0

        for celula in coluna:
            if celula.value is None:
                continue

            maior = max(
                maior,
                len(
                    str(
                        celula.value
                    )
                ),
            )

        ws.column_dimensions[
            letra
        ].width = min(
            maior + 3,
            40,
        )

    ws.freeze_panes = "A2"


def salvar_excel(
    df,
    caminho,
    fatal=True,
):
    caminho = Path(
        caminho
    )

    try:
        caminho.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        with pd.ExcelWriter(
            caminho,
            engine="openpyxl",
            date_format="DD/MM/YYYY",
            datetime_format="DD/MM/YYYY",
        ) as writer:
            df.to_excel(
                writer,
                index=False,
                sheet_name="calendario",
            )

            ajustar_excel(
                writer
            )

        return True

    except PermissionError:
        mensagem = (
            f"Não consegui salvar {caminho}. "
            f"Feche o arquivo no Excel e tente novamente."
        )

        if fatal:
            sys.exit(
                f"[ERRO] {mensagem}"
            )

        print(
            f"[ATENÇÃO] {mensagem}",
            file=sys.stderr,
        )

        return False


# ============================================================
# MAIN
# ============================================================

def main(
    headless=True,
    limite=None,
):
    print(
        "=" * 70
    )

    print(
        "COLETA DE HISTÓRICO DE JOGOS"
    )

    print(
        "=" * 70
    )

    print(
        f"Campeonato: {CAMPEONATO}"
    )

    print(
        f"País: {PAIS}"
    )

    print(
        f"URL: {COMPETICAO_URL}"
    )

    print()

    clubes_json = carregar_clubes()

    partes = []
    falhas = []
    sem_jogos = []

    sessao = Sessao(
        headless=headless,
    )

    try:
        print(
            f"Lendo temporadas: {COMPETICAO_URL}"
        )

        html = baixar_html(
            sessao,
            COMPETICAO_URL,
            EDICAO_SELECTOR,
        )

        edicoes = listar_edicoes(
            html
        )

        if not edicoes:
            raise RuntimeError(
                "Nenhuma edição encontrada."
            )

        print(
            f"Edições encontradas: {len(edicoes)}"
        )

        print(
            f"Mais recente: {edicoes[0]['temporada']}"
        )

        print(
            f"Mais antiga: {edicoes[-1]['temporada']}"
        )

        edicao_atual = next(
            (
                edicao
                for edicao in edicoes
                if edicao["atual"]
            ),
            edicoes[0],
        )

        salvar_link_temporada_atual(
            edicao_atual[
                "temporada"
            ],
            edicao_atual[
                "url_calendario"
            ],
        )

        if limite:
            edicoes = edicoes[
                :limite
            ]

        edicoes = list(
            reversed(
                edicoes
            )
        )

        print(
            f"\n{len(edicoes)} temporadas para processar.\n"
        )

        for i, edicao in enumerate(
            edicoes,
            start=1,
        ):
            temporada = edicao[
                "temporada"
            ]

            url_calendario = edicao[
                "url_calendario"
            ]

            print(
                f"[{i}/{len(edicoes)}] "
                f"Temporada {temporada}"
            )

            print(
                f"   {url_calendario}"
            )

            try:
                df_temporada = raspar_temporada(
                    sessao,
                    url_calendario,
                    temporada,
                    clubes_json,
                )

                validar_temporada(
                    df_temporada,
                    temporada,
                )

                if df_temporada.empty:
                    sem_jogos.append(
                        temporada
                    )
                else:
                    partes.append(
                        df_temporada
                    )

                if partes:
                    parcial = montar_saida(
                        partes
                    )

                    salvar_excel(
                        parcial,
                        OUTPUT_PATH,
                        fatal=False,
                    )

            except Exception as e:
                falhas.append(
                    temporada
                )

                print(
                    f"[FALHA] {temporada}: {resumo_erro(e)}",
                    file=sys.stderr,
                )

            pausa(
                1.5,
                3.0,
            )

    finally:
        sessao.encerrar()

    if not partes:
        sys.exit(
            "[ERRO] Nenhum jogo foi coletado."
        )

    df_final = montar_saida(
        partes
    )

    salvar_excel(
        df_final,
        OUTPUT_PATH,
        fatal=True,
    )

    print()
    print(
        "=" * 70
    )

    print(
        "FINALIZADO"
    )

    print(
        "=" * 70
    )

    print(
        f"Jogos finalizados: {len(df_final)}"
    )

    print(
        f"Temporadas com jogos: {len(partes)}"
    )

    print(
        f"Arquivo: {OUTPUT_PATH}"
    )

    if sem_jogos:
        print()
        print(
            "Temporadas sem jogos finalizados:"
        )

        for temporada in sem_jogos:
            print(
                f" - {temporada}"
            )

    if falhas:
        print()
        print(
            "Temporadas com falha:"
        )

        for temporada in falhas:
            print(
                f" - {temporada}"
            )


# ============================================================
# EXECUÇÃO
# ============================================================

if __name__ == "__main__":
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--headed",
        action="store_true",
    )

    parser.add_argument(
        "--limite",
        type=int,
        default=None,
    )

    args = parser.parse_args()

    main(
        headless=not args.headed,
        limite=args.limite,
    )

    try:
        import winsound

        winsound.Beep(
            1000,
            500,
        )

    except Exception:
        pass